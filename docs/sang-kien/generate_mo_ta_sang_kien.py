# -*- coding: utf-8 -*-
"""Generate filled innovation description Word doc for TNiCRM."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

DIR = Path(__file__).resolve().parent
OUT = DIR / "Mau_don_SK_TNiCRM_da_dien.docx"

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p

def add_body(doc, text):
    for para in text.strip().split("\n\n"):
        p = doc.add_paragraph(para.strip())
        p.paragraph_format.space_after = Pt(6)
        for run in p.runs:
            run.font.size = Pt(12)

def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MÔ TẢ SÁNG KIẾN\n")
    r.bold = True
    r.font.size = Pt(14)
    r2 = title.add_run(
        "Hệ thống quản lý quan hệ khách hàng và KPI bán hàng tích hợp CBS — TNiCRM"
    )
    r2.bold = True
    r2.font.size = Pt(13)

    sections = [
        (
            "1. Mục tiêu",
            """Xây dựng và triển khai hệ thống TNiCRM (CRM đa chi nhánh — multi-tenant) phục vụ đội kinh doanh B2B/B2C, tập trung hóa dữ liệu khách hàng — cơ hội — công việc — KPI tháng — doanh thu CBS; tự động hóa đối soát doanh thu và báo cáo tiến độ, giảm phụ thuộc file Excel rời rạc và thao tác thủ công giữa nhiều hệ thống.""",
        ),
        (
            "2. Lý do đề xuất giải pháp",
            """Trước khi có TNiCRM, đơn vị vận hành kinh doanh gặp các hạn chế:
• Dữ liệu khách hàng, deal và KPI nằm rải rác trên Excel, email và hệ thống CBS — khó đối soát “ai bán gì, thuộc kỳ nào”.
• Giao và theo dõi KPI tháng (doanh thu KH mới / hiện hữu, hợp đồng, tiếp cận) thủ công, dễ sai lệch khi chốt kỳ.
• Chuyển giao khách hàng giữa nhân viên thiếu luồng duyệt và nhật ký audit.
• Lãnh đạo không có dashboard thời gian thực theo phạm vi tổ chức (phòng/tổ/nhóm).

TNiCRM được phát triển nội bộ (Node.js + PostgreSQL + Next.js) nhằm khép kín vòng “Khách hàng → Deal → Thực hiện → KPI → Báo cáo”, đồng bộ doanh thu từ CBS và tôn trọng phân quyền theo vai trò (Admin, Giám đốc, Quản lý, Leader, Sales).""",
        ),
        (
            "3. Mục tiêu của giải pháp",
            """• Kinh tế — nâng năng suất đội sales và quản lý: giảm thời gian lập báo cáo, đối soát CBS, giao KPI; tăng độ chính xác số liệu thực hiện.
• Kỹ thuật — một nền tảng web tập trung, API chuẩn (Swagger), đa tenant, ma trận quyền cấu hình được.
• Xã hội / tổ chức — minh bạch chuyển giao khách, phê duyệt deal giá trị lớn, chốt kỳ KPI có hạn (ngày 5 tháng kế tiếp, giờ VN).
• Mở rộng — có thể nhân bản cho chi nhánh/đơn vị khác trong TCT qua mô hình tenant và Super Admin.""",
        ),
        (
            "4. Nội dung (tính mới và các hạng mục đã triển khai)",
            """4.1. Kiến trúc và nền tảng
• Backend: Node.js, Express, PostgreSQL, JWT; Frontend: Next.js (App Router), React.
• Đa tenant: mỗi đơn vị dữ liệu tách biệt theo tenant_id; đăng nhập theo domain/subdomain.
• Phân quyền RBAC: super_admin, admin, director, manager, leader, sales + ma trận quyền chi tiết (/permissions).
• Gói dịch vụ tenant: kiểm soát hạn sử dụng (xem khi hết hạn, chặn ghi khi quá hạn).

4.2. Module nghiệp vụ chính (đã vận hành)
(1) Khách hàng (/customers): CRUD, import Excel, tìm kiếm theo tên/công ty/SĐT/MST, lọc cột trên bảng, trạng thái hoạt động (active), phân loại NEW/EXISTING phục vụ KPI, lịch sử tương tác, chuyển giao hàng loạt.
(2) Cơ hội — Deal (/deals): Kanban pipeline (Tiếp cận → Tư vấn → Báo giá → Thương thảo → Thành công), loại MASS/Dự án, phê duyệt deal > 500 triệu, thùng rác khôi phục.
(3) Công việc — Task (/tasks): gán việc, hạn chót, liên kết KH/Deal, hoàn thành → ghi nhận tương tác.
(4) KPI tháng (/kpi): giao kế hoạch (doanh thu KH mới, KH hiện hữu, số hợp đồng, tiếp cận); thực hiện tự đồng bộ từ deal thắng + CBS; chốt kỳ; cấu hình cột hiển thị theo tenant; xuất Excel báo cáo.
(5) Doanh thu CBS (/cbs-revenue): đồng bộ đơn từ hệ thống CBS, gom theo số hợp đồng, gán người bán, duyệt, đồng bộ KPI; dashboard thẻ lọc nhanh.
(6) Tổ chức & nhân sự: cây phòng ban (/organization-structure), nhân viên import Excel (/employees), chuyển giao & duyệt (/transfers, /transfer-approvals), nhật ký (/transfer-logs).
(7) Dashboard (/dashboard): thống kê, tiến độ KPI tổng, bảng xếp hạng KPI tháng, biểu đồ.

4.3. Điểm mới nổi bật so với CRM thông thường
• Tích hợp hai nguồn thực hiện KPI: Deal “Thành công” + đơn CBS đã duyệt (gom HĐ, 1 KPI/HĐ).
• Tiếp cận KPI = deal trong kỳ + khách CBS đã duyệt (loại trùng).
• Luồng chuyển giao khách có phê duyệt và audit.
• Cấu hình hiển thị cột KPI lưu server theo tenant (không phụ thuộc từng máy).

[Hình ảnh — Sơ đồ tổng thể kiến trúc: vẽ khối Frontend (Next.js) ↔ API (Express) ↔ PostgreSQL; nhánh tích hợp CBS. Đặt Phụ lục 1.]

[Hình ảnh trang /login — Màn hình đăng nhập theo domain tenant]

[Hình ảnh trang /dashboard — Dashboard tổng quan: thẻ thống kê, biểu đồ KPI, bộ lọc thời gian]

[Hình ảnh trang /customers — Danh sách khách hàng: tìm kiếm, lọc cột, MST, trạng thái]

[Hình ảnh trang /deals — Kanban pipeline cơ hội bán hàng]

[Hình ảnh trang /kpi — Bảng KPI tháng: STT, Nhân viên, Địa bàn, Hợp đồng, DT Tổng, DT mới, DT HH, Tiếp cận (KH/TH)]

[Hình ảnh trang /cbs-revenue — Doanh thu CBS: thẻ thống kê có thể click lọc, bảng gom theo hợp đồng]

[Hình ảnh trang /organization-structure — Cơ cấu tổ chức dạng cây]

[Hình ảnh trang /permissions — Ma trận phân quyền theo vai trò]""",
        ),
        (
            "5. Tính mới của sáng kiến (cam đoan theo quy định)",
            """• Giải pháp do tác giả/đơn vị tự xây dựng cho nghiệp vụ kinh doanh viễn thông/dịch vụ số, không sao chép nguyên xi sản phẩm CRM thương mại.
• Tích hợp KPI–CBS–Deal trong một luồng dữ liệu thống nhất là đặc thù triển khai tại đơn vị, chưa thấy quy trình/tiêu chuẩn bắt buộc tương đương nội bộ trước đây.
• Mã nguồn và cấu hình nằm trong kiểm soát đơn vị; có thể mở rộng tenant mà không phụ thuộc license CRM bên thứ ba.
• Chưa công bố rộng rãi ngoài phạm vi triển khai nội bộ (điền thêm ngày đưa vào vận hành thực tế nếu có quyết định).""",
        ),
        (
            "6. Đã được áp dụng hoặc áp dụng thử",
            """Hệ thống đã triển khai môi trường phát triển/vận hành thử nghiệm với đầy đủ module: đăng nhập đa tenant, quản lý khách hàng & deal, task, KPI tháng (giao kế hoạch, đồng bộ thực hiện, xuất báo cáo), đồng bộ CBS, dashboard, phân quyền, chuyển giao khách.

Các tính năng gần đây đã hoàn thiện: cột STT và thứ tự cột KPI; thuật ngữ “kế hoạch” (KH) thay cho chỉ tiêu; một kế hoạch số lượng hợp đồng; gom đơn CBS theo contract_no; tìm kiếm khách hàng nâng cao; trạng thái active khách hàng.

(Điền tên đơn vị, thời gian go-live, số user thực tế khi nộp hồ sơ chính thức.)""",
        ),
        (
            "7. Kết quả áp dụng",
            """7.1. Đối tượng / quy trình TRƯỚC khi áp dụng
• Sales cập nhật khách và cơ hội trên file hoặc công cụ rời; quản lý tổng hợp KPI cuối tháng bằng Excel.
• Doanh thu CBS phải đối chiếu thủ công với danh sách nhân viên — dễ trùng, thiếu người bán, chậm chốt số.
• Chuyển khách giữa nhân viên qua trao đổi informal, khó truy vết.
• Lãnh đạo thiếu một màn hình thống nhất theo kỳ và theo phòng ban.

7.2. Đối tượng / quy trình SAU khi áp dụng
• Một cổng web: nhập liệu tại chỗ, phân quyền đúng người đúng việc.
• KPI tháng: giao kế hoạch trên /kpi; thực hiện cập nhật từ deal + CBS; chốt kỳ tự động sau hạn ngày 5 (VN).
• CBS: đồng bộ đơn, gom hợp đồng, gán seller, một đơn vị KPI hợp đồng.
• Dashboard và Excel export phục vụ họp giao ban đầu/cuối tháng.

7.3. Minh chứng đề nghị đính kèm
[Hình ảnh / Bảng — So sánh thời gian lập báo cáo KPI: trước (giờ/người) vs sau (phút)]

[Hình ảnh — File Excel xuất từ /kpi: Bao_cao_KPI_thang_YYYY-MM.xlsx]

[Hình ảnh — Màn hình /cbs-revenue sau khi gán người bán và trạng thái đã duyệt]

[Email / Biên bản — Xác nhận triển khai hoặc kết quả pilot từ lãnh đạo đơn vị — người nộp điền]

(Điền số liệu cụ thể: số NV dùng, số KH, số deal/tháng, số đơn CBS đồng bộ — khi có số liệu vận hành.)""",
        ),
        (
            "8. Các điều kiện cần thiết để áp dụng sáng kiến",
            """• Hạ tầng: máy chủ hoặc cloud chạy Node.js, PostgreSQL; truy cập HTTPS nội bộ hoặc VPN.
• Tài khoản tích hợp CBS (username/password/unit_id) do Admin cấu hình tại /cbs-revenue.
• Đào tạo ngắn theo vai trò (Sales / Quản lý / Admin); tài liệu: docs/HUONG_DAN_SU_DUNG.md trong mã nguồn.
• Quy ước nghiệp vụ: quy trình chuyển giao khách, ngưỡng phê duyệt deal, lịch giao KPI đầu tháng.
• Bảo trì: cập nhật migration DB khi nâng phiên bản; sao lưu PostgreSQL định kỳ.""",
        ),
        (
            "9. Đánh giá lợi ích thu được",
            """9.1. Về kinh tế
• Giảm chi phí ẩn của nhập liệu trùng và báo cáo thủ công (ước tính điền: X giờ/người/tháng × Y nhân sự × chi phí giờ công).
• Tăng độ tin cậy số liệu KPI → quyết định giao chỉ tiêu kỳ sau chính xác hơn.
• Tránh bỏ sót doanh thu CBS không gán seller.

9.2. Về thời gian
• Tổng hợp KPI tháng: từ thu thập nhiều nguồn → một lần mở /kpi hoặc xuất Excel.
• Đối soát CBS: lọc theo thẻ trạng thái trên /cbs-revenue thay vì lọc file thủ công.
• Tìm khách: search đa trường (tên, MST, SĐT) thay vì lật nhiều sheet.

9.3. Về nhân lực
• Không cần thêm nhân sự chuyên “tổng hợp Excel” full-time; quản lý tập trung giám sát trên dashboard.
• Sales tự theo dõi % đạt kế hoạch (chế độ Số / % trên KPI).
• Admin tenant một lần cấu hình tổ chức + quyền, tái sử dụng cho nhân viên mới (import Excel).

(Phương pháp tính tiền làm lợi — điền khi có số liệu: Δ thời gian × đơn giá giờ công + giá trị tránh sai sót KPI.)""",
        ),
        (
            "10. Khả năng áp dụng",
            """• Cấp đơn vị: triển khai ngay cho toàn bộ phòng kinh doanh/dịch vụ đang dùng CBS và quản lý KH bằng Excel.
• Cấp TCT: nhân rộng qua mô hình multi-tenant (mỗi chi nhánh một tenant), Super Admin quản lý tại /system/tenants.
• Điều kiện: có quy trình bán hàng qua pipeline deal và có nguồn doanh thu CBS (hoặc tương đương có thể tích hợp API).
• Không yêu cầu thay thế CBS — CRM bổ sung lớp quản trị KH, KPI và người bán.""",
        ),
    ]

    for heading, body in sections:
        add_heading(doc, heading, level=2)
        add_body(doc, body)

    doc.add_page_break()
    add_heading(doc, "PHỤ LỤC — Danh sách hình ảnh đề nghị chèn", level=2)
    add_body(
        doc,
        """| STT | Mô tả hình | Trang / chức năng |
|-----|------------|------------------|
| 1 | Sơ đồ kiến trúc hệ thống | Vẽ tay hoặc draw.io |
| 2 | Đăng nhập | /login |
| 3 | Dashboard | /dashboard |
| 4 | Danh sách khách hàng | /customers |
| 5 | Chi tiết khách + timeline | /customers/[id] |
| 6 | Kanban Deal | /deals |
| 7 | KPI tháng (bảng đầy đủ cột) | /kpi |
| 8 | Cấu hình cột KPI (Admin) | /kpi — nút Cấu hình cột |
| 9 | Doanh thu CBS + thẻ lọc | /cbs-revenue |
| 10 | Cơ cấu tổ chức | /organization-structure |
| 11 | Ma trận phân quyền | /permissions |
| 12 | Chuyển giao / Duyệt chuyển giao | /transfers, /transfer-approvals |
| 13 | Xuất Excel KPI | File tải từ /kpi |
| 14 | Swagger API | /api-docs (backend) |

Ghi chú chụp màn hình: ẩn/bôi mờ dữ liệu nhạy cảm (SĐT, MST thật, doanh thu chi tiết nếu cần).""",
    )

    doc.save(OUT)
    print("OK")

if __name__ == "__main__":
    main()
